'''실제 엑셀 읽기, 파일 변경 반영, 오류 거부와 원본 보존 검증.'''

import os
import shutil
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from data.excel_loader import (
  DUMMY_DIR,
  FILES,
  DataValidationError,
  load_tables,
  validate_tables,
)
from data.sample_data import load_dashboard_snapshot
from tests.excel_fixtures import change_cells


class ExcelLoaderTest(unittest.TestCase):
  @classmethod
  def setUpClass(cls):
    cls.tables = load_tables()

  def test_preserves_dictionary_and_pattern_source_rows(self):
    for name in ('sipri_dictionary', 'nato_dictionary', 'usage_patterns'):
      filename, sheet = FILES[name]
      expected = pd.read_excel(DUMMY_DIR / filename, sheet_name=sheet)
      pd.testing.assert_frame_equal(self.tables[name], expected)
    self.assertEqual(self.tables['usage_patterns']['trans_form'].isna().sum(), 31)
    self.assertEqual(self.tables['sipri_dictionary'].duplicated('wp_name').sum(), 4)

  def test_new_workbooks_have_typed_dates_and_expected_sheets(self):
    for name in ('articles', 'result'):
      filename, sheet = FILES[name]
      book = load_workbook(DUMMY_DIR / filename, read_only=True)
      try:
        self.assertEqual(book.sheetnames, [sheet])
        self.assertIsInstance(book[sheet]['A2'].value, int)
        if name == 'articles':
          self.assertTrue(book[sheet]['C2'].is_date)
      finally:
        book.close()

  def test_rejects_invalid_columns_keys_relations_dates_codes_and_evidence(self):
    cases = [
      ('categories', 'kind', 'weapon', 'kind'),
      ('categories', 'category_name', 123, '문자열'),
      ('articles', 'conflict_id', 'missing', '참조'),
      ('articles', 'published_date', 'invalid', '날짜'),
      ('articles', 'published_date', 12345, '날짜'),
      ('articles', 'article_id', 1.5, '정수'),
      ('result', 'category_id', 'missing', '참조'),
      ('result', 'usage_code', 3, '정수'),
      ('conflicts', 'latitude', 100, '범위'),
      ('conflicts', 'color', 'red', '형식'),
      ('sipri_dictionary', 'category_id', 'T0101', '유형'),
    ]
    for name, column, value, message in cases:
      with self.subTest(table=name, column=column, value=value):
        tables = {key: frame.copy(deep=True) for key, frame in self.tables.items()}
        tables[name][column] = tables[name][column].astype(object)
        tables[name].loc[0, column] = value
        with self.assertRaisesRegex(DataValidationError, message):
          validate_tables(tables)
    for mode, message in [
      ('column', '필수 컬럼'),
      ('duplicate', '중복'),
      ('evidence', '근거|evidence'),
    ]:
      with self.subTest(mode=mode):
        tables = {key: frame.copy(deep=True) for key, frame in self.tables.items()}
        if mode == 'column':
          tables['articles'] = tables['articles'].drop(columns='title')
        elif mode == 'duplicate':
          tables['result'] = pd.concat(
            [tables['result'], tables['result'].iloc[:1]], ignore_index=True
          )
        else:
          tables['result'].loc[0, ['usage_code', 'evidence_sentence']] = [1, '  ']
        with self.assertRaisesRegex(DataValidationError, message):
          validate_tables(tables)

  def test_empty_result_and_article_files_are_valid(self):
    tables = {key: frame.copy() for key, frame in self.tables.items()}
    tables['result'] = tables['result'].iloc[:0]
    validate_tables(tables)
    tables['articles'] = tables['articles'].iloc[:0]
    validate_tables(tables)

  def test_file_changes_refresh_metadata_and_counts_even_with_same_mtime(self):
    with tempfile.TemporaryDirectory() as folder:
      folder = Path(folder)
      for filename, _ in FILES.values():
        shutil.copy2(DUMMY_DIR / filename, folder / filename)
      before = load_dashboard_snapshot(folder)
      category_file = folder / 'categories.xlsm'
      category_stat = category_file.stat()
      change_cells(category_file, {'C2': '수정된 기타'})
      os.utime(category_file, ns=(category_stat.st_atime_ns, category_stat.st_mtime_ns))
      result_file = folder / 'result_example.xlsx'
      # 첫 결과는 사용 외 언급이며 이 기사의 결과는 한 건뿐이다.
      change_cells(result_file, {'C2': 1})
      after = load_dashboard_snapshot(folder)
      self.assertIn('수정된 기타', after['settings']['categories']['기술'])
      self.assertNotIn('기타', after['settings']['categories']['기술'])

      def total(snapshot):
        return int(
          snapshot['article_daily'].query('kind == "전체"')['usage_articles'].sum()
        )

      self.assertEqual(total(after), total(before) + 1)
      # 잘못된 최신 파일은 이전 성공 캐시로 대체하지 않는다.
      change_cells(result_file, {'B2': 'missing'})
      with self.assertRaisesRegex(DataValidationError, 'result_example.xlsx.*참조'):
        load_dashboard_snapshot(folder)
      result_file.unlink()
      with self.assertRaisesRegex(DataValidationError, 'result_example.xlsx'):
        load_dashboard_snapshot(folder)

  def test_unreadable_workbook_reports_filename(self):
    with tempfile.TemporaryDirectory() as folder:
      folder = Path(folder)
      for filename, _ in FILES.values():
        shutil.copy2(DUMMY_DIR / filename, folder / filename)
      (folder / 'categories.xlsm').write_bytes(b'not an Excel workbook')
      with self.assertRaisesRegex(DataValidationError, 'categories.xlsm'):
        load_tables(folder)


if __name__ == '__main__':
  unittest.main()
